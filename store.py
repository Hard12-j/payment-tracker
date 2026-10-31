"""Storage for BOTH the public form and the admin app.

Two interchangeable backends (choose with STORAGE_BACKEND):
  "sheets"  Google Sheets  -> use this when deployed (survives restarts, no local disk needed)
  "excel"   local payments.xlsx -> fine on your own PC / a server with a permanent disk

Tab "Payments"          one row per payment.
                        Receipt No = auto RCP-0001... for admin/cash entries,
                                   = the TRANSACTION ID for online (public form) payments.
                        Receipt Link = Drive link of the PDF receipt (admin entries) or of the
                                       payer's payment screenshot (online payments).
Tab "Rejected Attempts" audit log of submissions the public form rejected (admin only).

All writes go through one lock, so two people submitting at once cannot clash.
"""
import io
import os
import re
import threading
import time
from datetime import datetime
from pathlib import Path

import pandas as pd

import verify
from settings import (EXCEL_FILE, RECEIPT_PREFIX, SHEETS_NAME, SHEETS_SPREADSHEET_ID,
                      STORAGE_BACKEND)

SHEET = "Payments"
ATTEMPT_SHEET = "Rejected Attempts"

SOURCE_ADMIN = "Admin"
SOURCE_ONLINE = "Online"

# First 10 columns are exactly the original gayatri-nagar layout (old files keep working).
HEADERS = [
    "Receipt No", "Date", "Name", "Plot No", "Amount",
    "Payment Mode", "Email", "Phone", "Remarks", "Receipt Link",
    "Source", "Bank Status", "Screenshot Hash",
]
COL_WIDTHS = {
    "Receipt No": 22, "Date": 20, "Name": 26, "Plot No": 12, "Amount": 14,
    "Payment Mode": 14, "Email": 28, "Phone": 16, "Remarks": 30, "Receipt Link": 40,
    "Source": 10, "Bank Status": 14, "Screenshot Hash": 20,
}
ATTEMPT_HEADERS = ["Date", "Name", "Plot No", "Amount", "Transaction ID", "Reason"]

_lock = threading.RLock()


class DuplicateError(Exception):
    """Raised when a transaction ID or screenshot is already recorded."""


def now_str() -> str:
    return datetime.now().strftime("%d-%m-%Y %H:%M")


def _col_letter(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _cell(row, i):
    return row[i] if i < len(row) and row[i] is not None else ""


# ==========================================================================
# Backends. Tiny interface: rows are lists of strings, row/col numbers are 1-based.
#   ensure_sheet(title, headers)   create the tab if missing; write headers if it is empty
#   read(title, fresh=False)       all rows incl. header ([] if the tab does not exist)
#   append(title, row)             add one row at the bottom
#   update_cells(title, cells)     cells = [(row, col, value), ...]
# ==========================================================================
class ExcelBackend:
    cache_schema = False        # cheap to re-check, and the file may be replaced by hand

    def __init__(self, path, widths=None):
        self.path = Path(path)
        self.widths = widths or {}
        self._schema_ok = False

    def _save(self, wb):
        tmp = self.path.with_name(f".{self.path.stem}.tmp.xlsx")
        wb.save(tmp)
        os.replace(tmp, self.path)       # atomic: a crash never leaves a half-written file

    def ensure_sheet(self, title, headers):
        from openpyxl import Workbook, load_workbook
        with _lock:
            if not self.path.exists():
                wb = Workbook()
                ws = wb.active
                ws.title = title
                self._write_headers(ws, headers)
                self._save(wb)
                return
            wb = load_workbook(self.path)
            try:
                ws = wb[title] if title in wb.sheetnames else wb.create_sheet(title)
                if ws.max_row == 1 and ws.max_column == 1 and ws.cell(1, 1).value is None:
                    self._write_headers(ws, headers)   # empty/new tab -> write the header row
                    self._save(wb)
            finally:
                wb.close()

    def _write_headers(self, ws, headers):
        # explicit cells (not ws.append): append() would skip row 1 once A1 has been touched
        for i, h in enumerate(headers, start=1):
            ws.cell(row=1, column=i, value=h)
        self._set_widths(ws, headers)

    def _set_widths(self, ws, headers):
        for i, h in enumerate(headers, start=1):
            if h in self.widths:
                ws.column_dimensions[_col_letter(i)].width = self.widths[h]

    def read(self, title, fresh=False):
        from openpyxl import load_workbook
        if not self.path.exists():
            return []
        wb = load_workbook(self.path, read_only=True)
        try:
            if title not in wb.sheetnames:
                return []
            rows = [["" if v is None else str(v) for v in r]
                    for r in wb[title].iter_rows(values_only=True)]
        finally:
            wb.close()
        while rows and not any(rows[-1]):
            rows.pop()
        return rows

    def append(self, title, row):
        from openpyxl import load_workbook
        with _lock:
            wb = load_workbook(self.path)
            try:
                wb[title].append(list(row))
                self._save(wb)
            finally:
                wb.close()

    def update_cells(self, title, cells):
        from openpyxl import load_workbook
        if not cells:
            return
        with _lock:
            wb = load_workbook(self.path)
            try:
                ws = wb[title]
                for r, c, v in cells:
                    cell = ws.cell(row=r, column=c, value=v)
                    if r == 1 and v in self.widths:
                        ws.column_dimensions[cell.column_letter].width = self.widths[v]
                self._save(wb)
            finally:
                wb.close()

    def url(self):
        return ""


class SheetsBackend:
    """Google Sheets via the Sheets API. Values are written RAW so long transaction IDs and
    amounts stay exactly as text (no scientific notation, no lost leading zeros)."""
    cache_schema = True          # schema is checked once per process (saves API calls)
    CACHE_TTL = 4                # seconds; admin screens re-run often, quota is limited

    def __init__(self, spreadsheet_id="", title="Gayatri Nagar Payments"):
        self._id = spreadsheet_id
        self.title = title
        self._cache = {}
        self._schema_ok = False

    # -- helpers
    @staticmethod
    def _a1(title):
        return "'" + title.replace("'", "''") + "'"

    def _svc(self):
        import drive_uploader
        return drive_uploader.get_sheets_service()

    @property
    def spreadsheet_id(self):
        if not self._id:
            import drive_uploader
            self._id = drive_uploader.find_or_create_spreadsheet(self.title, SHEET)
        return self._id

    def url(self):
        return f"https://docs.google.com/spreadsheets/d/{self.spreadsheet_id}"

    @staticmethod
    def _lock_api():
        import drive_uploader
        return drive_uploader.API_LOCK

    # -- interface
    def ensure_sheet(self, title, headers):
        with self._lock_api():
            svc, sid = self._svc(), self.spreadsheet_id
            meta = svc.spreadsheets().get(spreadsheetId=sid, fields="sheets.properties.title").execute()
            titles = [s["properties"]["title"] for s in meta.get("sheets", [])]
            if title not in titles:
                svc.spreadsheets().batchUpdate(
                    spreadsheetId=sid,
                    body={"requests": [{"addSheet": {"properties": {"title": title}}}]}).execute()
            first = svc.spreadsheets().values().get(
                spreadsheetId=sid, range=f"{self._a1(title)}!1:1").execute()
            if not first.get("values"):
                svc.spreadsheets().values().update(
                    spreadsheetId=sid, range=f"{self._a1(title)}!A1", valueInputOption="RAW",
                    body={"values": [list(headers)]}).execute()
            self._cache.pop(title, None)

    def read(self, title, fresh=False):
        hit = self._cache.get(title)
        if hit and not fresh and time.time() - hit[0] < self.CACHE_TTL:
            return [list(r) for r in hit[1]]
        with self._lock_api():
            try:
                res = self._svc().spreadsheets().values().get(
                    spreadsheetId=self.spreadsheet_id, range=self._a1(title),
                    valueRenderOption="FORMATTED_VALUE").execute()
            except Exception as e:      # tab does not exist yet
                if "Unable to parse range" in str(e):
                    return []
                raise
            rows = [[str(v) for v in r] for r in res.get("values", [])]
        self._cache[title] = (time.time(), rows)
        return [list(r) for r in rows]

    def append(self, title, row):
        with self._lock_api():
            self._svc().spreadsheets().values().append(
                spreadsheetId=self.spreadsheet_id, range=f"{self._a1(title)}!A1",
                valueInputOption="RAW", insertDataOption="INSERT_ROWS",
                body={"values": [[str(v) for v in row]]}).execute()
            self._cache.pop(title, None)

    def update_cells(self, title, cells):
        if not cells:
            return
        with self._lock_api():
            svc, sid = self._svc(), self.spreadsheet_id
            for i in range(0, len(cells), 400):
                data = [{"range": f"{self._a1(title)}!{_col_letter(c)}{r}", "values": [[str(v)]]}
                        for r, c, v in cells[i:i + 400]]
                svc.spreadsheets().values().batchUpdate(
                    spreadsheetId=sid, body={"valueInputOption": "RAW", "data": data}).execute()
            self._cache.pop(title, None)


# ==========================================================================
# Choosing the backend
# ==========================================================================
_default_backend = None


def get_backend():
    global _default_backend
    if _default_backend is None:
        if STORAGE_BACKEND == "sheets":
            _default_backend = SheetsBackend(SHEETS_SPREADSHEET_ID, SHEETS_NAME)
        else:
            _default_backend = ExcelBackend(EXCEL_FILE, COL_WIDTHS)
    return _default_backend


def _be(target=None):
    """target: None (configured backend), a file path (Excel file), or a backend object.
    The parameter is called file_path in the public functions for backwards compatibility."""
    if target is None:
        return get_backend()
    if isinstance(target, (ExcelBackend, SheetsBackend)) or hasattr(target, "ensure_sheet"):
        return target
    return ExcelBackend(Path(target), COL_WIDTHS)


def storage_url(file_path=None) -> str:
    """Link to the Google Sheet (empty for local Excel)."""
    try:
        return _be(file_path).url()
    except Exception:
        return ""


# ==========================================================================
# Schema: create tab / add missing columns / label pre-merge rows as Admin
# ==========================================================================
def ensure_schema(be, force=False):
    if be.cache_schema and be._schema_ok and not force:
        return
    with _lock:
        rows = be.read(SHEET, fresh=True)
        if not rows:
            be.ensure_sheet(SHEET, HEADERS)
        else:
            headers = rows[0]
            missing = [h for h in HEADERS if h not in headers]
            if missing:
                start = len([h for h in headers if h]) + 1
                cells = [(1, start + i, h) for i, h in enumerate(missing)]
                if "Source" in missing:
                    # rows from before the merge were all entered by the admin
                    new_headers = [h for h in headers if h] + missing
                    src_col = new_headers.index("Source") + 1
                    for r, row in enumerate(rows[1:], start=2):
                        if str(_cell(row, 0)).strip():
                            cells.append((r, src_col, SOURCE_ADMIN))
                be.update_cells(SHEET, cells)
        be._schema_ok = True


def migrate_excel_if_needed(file_path=None):
    """Upgrade an existing file/sheet to the current columns (no-op if nothing to do)."""
    be = _be(file_path)
    if isinstance(be, ExcelBackend) and not be.path.exists():
        return
    ensure_schema(be, force=True)


def _ensure_workbook(file_path=None):
    ensure_schema(_be(file_path), force=True)


def _headers_and_rows(be):
    rows = be.read(SHEET, fresh=True)
    return (rows[0] if rows else list(HEADERS)), rows[1:]


# ==========================================================================
# Receipt numbers (cash / admin entries)
# ==========================================================================
def next_receipt_no(file_path=None) -> str:
    """Next RCP-000N. Only values that look like <PREFIX>-<digits> are counted, so an online
    payment whose Receipt No is a long transaction ID can never disturb the sequence."""
    be = _be(file_path)
    try:
        ensure_schema(be)
        rows = be.read(SHEET)
    except PermissionError:
        return f"{RECEIPT_PREFIX}-????"
    pattern = re.compile(rf"^{re.escape(RECEIPT_PREFIX)}-(\d+)$")
    highest = 0
    for row in rows[1:]:
        m = pattern.match(str(_cell(row, 0)).strip())
        if m:
            highest = max(highest, int(m.group(1)))
    return f"{RECEIPT_PREFIX}-{highest + 1:04d}"


# ==========================================================================
# Writing
# ==========================================================================
def _record_to_row(record: dict, headers):
    return ["" if (h is None or record.get(h) is None) else str(record.get(h)) for h in headers]


def save_payment(record: dict, file_path=None):
    """Append one row. Columns are matched by header name, so column order does not matter."""
    be = _be(file_path)
    ensure_schema(be)
    with _lock:
        headers, _ = _headers_and_rows(be)
        be.append(SHEET, _record_to_row(record, headers))


def add_admin_payment(record: dict, file_path=None):
    record = dict(record)
    record.setdefault("Source", SOURCE_ADMIN)
    save_payment(record, file_path)


def _find_duplicate_in(headers, rows, txn_key, img_hash):
    i_no = headers.index("Receipt No") if "Receipt No" in headers else 0
    i_hash = headers.index("Screenshot Hash") if "Screenshot Hash" in headers else None
    for row in rows:
        no = _cell(row, i_no)
        if no and verify.normalize_id(no) == txn_key:
            return "This transaction ID has already been used"
        if i_hash is not None and img_hash and _cell(row, i_hash) == img_hash:
            return "This screenshot has already been submitted"
    return None


def find_duplicate(txn_key: str, img_hash: str, file_path=None):
    """Reason string if the txn id (normalised) or screenshot is already recorded, else None."""
    be = _be(file_path)
    ensure_schema(be)
    with _lock:
        headers, rows = _headers_and_rows(be)
        return _find_duplicate_in(headers, rows, txn_key, img_hash)


def add_online_payment(name, plot_no, amount, txn_id, img_hash, receipt_link="", remarks="",
                       file_path=None) -> dict:
    """Save a public-form payment. Receipt No = transaction ID; Receipt Link = Drive link of the
    payer's screenshot. The duplicate check and the write happen under one lock (no race).
    Raises DuplicateError with the reason if it is already recorded."""
    be = _be(file_path)
    ensure_schema(be)
    receipt_no = verify.clean_txn_id(txn_id)
    key = verify.normalize_id(receipt_no)
    with _lock:
        headers, rows = _headers_and_rows(be)
        reason = _find_duplicate_in(headers, rows, key, img_hash)
        if reason:
            raise DuplicateError(reason)
        record = {
            "Receipt No": receipt_no,
            "Date": now_str(),
            "Name": name,
            "Plot No": plot_no,
            "Amount": f"{float(amount):.2f}",
            "Payment Mode": "Online",
            "Remarks": remarks,
            "Receipt Link": receipt_link,
            "Source": SOURCE_ONLINE,
            "Bank Status": "Unconfirmed",
            "Screenshot Hash": img_hash,
        }
        be.append(SHEET, _record_to_row(record, headers))
    return record


def update_receipt_link(receipt_no: str, link: str, file_path=None):
    """Find the row by Receipt No and write the Drive link into 'Receipt Link'."""
    be = _be(file_path)
    ensure_schema(be)
    with _lock:
        headers, rows = _headers_and_rows(be)
        i_no, i_link = headers.index("Receipt No"), headers.index("Receipt Link")
        for r, row in enumerate(rows, start=2):
            if str(_cell(row, i_no)).strip() == str(receipt_no).strip():
                be.update_cells(SHEET, [(r, i_link + 1, link)])
                return


def mark_bank_confirmed(receipt_nos, file_path=None) -> int:
    """Bank Status = Confirmed for ONLINE rows whose Receipt No (transaction ID) is given.
    Returns how many rows changed."""
    keys = {verify.normalize_id(n) for n in receipt_nos}
    if not keys:
        return 0
    be = _be(file_path)
    ensure_schema(be)
    with _lock:
        headers, rows = _headers_and_rows(be)
        i_no, i_src, i_bank = (headers.index(h) for h in ("Receipt No", "Source", "Bank Status"))
        cells = []
        for r, row in enumerate(rows, start=2):
            if _cell(row, i_src) != SOURCE_ONLINE or _cell(row, i_bank) == "Confirmed":
                continue
            if verify.normalize_id(_cell(row, i_no)) in keys:
                cells.append((r, i_bank + 1, "Confirmed"))
        be.update_cells(SHEET, cells)
        return len(cells)


def log_attempt(name, plot_no, amount, txn_id, reason, file_path=None):
    """Audit log of rejected public submissions (visible to admin only)."""
    be = _be(file_path)
    ensure_schema(be)
    with _lock:
        be.ensure_sheet(ATTEMPT_SHEET, ATTEMPT_HEADERS)
        be.append(ATTEMPT_SHEET, [now_str(), str(name), str(plot_no), str(amount),
                                  str(txn_id), reason])


# ==========================================================================
# Reading (admin)
# ==========================================================================
def _to_df(rows, default_headers):
    if not rows:
        return pd.DataFrame(columns=default_headers)
    headers = [h if h else f"col{i + 1}" for i, h in enumerate(rows[0])]
    width = len(headers)
    body = [[str(_cell(r, i)) for i in range(width)] for r in rows[1:]]
    return pd.DataFrame(body, columns=headers)


def load_payments(file_path=None) -> pd.DataFrame:
    be = _be(file_path)
    try:
        ensure_schema(be)
        return _to_df(be.read(SHEET), HEADERS)
    except PermissionError:
        return pd.DataFrame(columns=HEADERS)


def load_attempts(file_path=None) -> pd.DataFrame:
    be = _be(file_path)
    try:
        ensure_schema(be)
        return _to_df(be.read(ATTEMPT_SHEET), ATTEMPT_HEADERS)
    except PermissionError:
        return pd.DataFrame(columns=ATTEMPT_HEADERS)


def export_xlsx_bytes(file_path=None) -> bytes:
    """Both tabs as an .xlsx file (for the admin 'Download Excel' button)."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        load_payments(file_path).to_excel(w, sheet_name=SHEET, index=False)
        load_attempts(file_path).to_excel(w, sheet_name=ATTEMPT_SHEET, index=False)
    return buf.getvalue()
