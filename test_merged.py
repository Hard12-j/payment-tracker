"""Offline tests for the merged app (no OCR model, no browser, no Google needed).
Run:  python test_merged.py
"""
import io
import shutil
import tempfile
import threading
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

import bank_check
import drive_uploader
import receipts
import store
import verify


# --------------------------------------------------------------------------
# In-memory imitation of the Google Sheets API (only the calls store.py makes)
# --------------------------------------------------------------------------
import re as _re


class _Exec:
    def __init__(self, fn): self.fn = fn
    def execute(self): return self.fn()


class FakeSheetsService:
    def __init__(self):
        self.tabs = {}            # title -> list of rows (lists of str)
        self.calls = 0

    def _split(self, rng):
        m = _re.match(r"^'((?:[^']|'')*)'(?:!(.*))?$", rng)
        return m.group(1).replace("''", "'"), m.group(2)

    def spreadsheets(self): return self

    def get(self, spreadsheetId, fields=None):
        return _Exec(lambda: {"sheets": [{"properties": {"title": t}} for t in self.tabs]})

    def batchUpdate(self, spreadsheetId, body):
        def run():
            for req in body["requests"]:
                self.tabs.setdefault(req["addSheet"]["properties"]["title"], [])
            return {}
        return _Exec(run)

    def values(self): return _Values(self)


class _Values:
    def __init__(self, svc): self.svc = svc

    def get(self, spreadsheetId, range, valueRenderOption=None):
        def run():
            self.svc.calls += 1
            title, part = self.svc._split(range)
            if title not in self.svc.tabs:
                raise Exception(f"Unable to parse range: {range}")
            rows = self.svc.tabs[title]
            if part == "1:1":
                rows = rows[:1]
            out = []
            for r in rows:
                r = list(r)
                while r and r[-1] == "":
                    r.pop()                      # the real API drops trailing empty cells
                out.append(r)
            while out and out[-1] == []:
                out.pop()
            return {"values": out} if out else {}
        return _Exec(run)

    def update(self, spreadsheetId, range, valueInputOption, body):
        def run():
            title, part = self.svc._split(range)
            assert part == "A1" and valueInputOption == "RAW"
            self.svc.tabs[title] = [list(r) for r in body["values"]]
            return {}
        return _Exec(run)

    def append(self, spreadsheetId, range, valueInputOption, insertDataOption, body):
        def run():
            title, _ = self.svc._split(range)
            assert valueInputOption == "RAW" and insertDataOption == "INSERT_ROWS"
            assert all(isinstance(v, str) for r in body["values"] for v in r)   # RAW text only
            self.svc.tabs[title].extend(list(r) for r in body["values"])
            return {}
        return _Exec(run)

    def batchUpdate(self, spreadsheetId, body):
        def run():
            assert body["valueInputOption"] == "RAW"
            for item in body["data"]:
                title, cell = self.svc._split(item["range"])
                col_s, row_s = _re.match(r"([A-Z]+)(\d+)", cell).groups()
                col = 0
                for ch in col_s:
                    col = col * 26 + ord(ch) - 64
                row = int(row_s)
                rows = self.svc.tabs[title]
                while len(rows) < row:
                    rows.append([])
                while len(rows[row - 1]) < col:
                    rows[row - 1].append("")
                rows[row - 1][col - 1] = item["values"][0][0]
            return {}
        return _Exec(run)


def shot_text(status="Payment Successful", amount="Rs 5,000", utr="412345678901"):
    return f"{status}\n{amount}\nPaid to shop@okbank\nUPI Transaction ID\n{utr}\n31 Oct 2026, 10:42 AM"


class TestVerifyNoWordRejection(unittest.TestCase):
    def test_valid(self):
        self.assertEqual(verify.run_checks(shot_text(), "412345678901", 5000), [])

    def test_status_words_never_reject(self):
        for word in ("Payment Failed", "Pending", "Processing", "Declined", "", "Transaction Failed"):
            with self.subTest(word=word):
                self.assertEqual(verify.run_checks(shot_text(status=word), "412345678901", 5000), [])

    def test_no_status_function_left(self):
        self.assertFalse(hasattr(verify, "check_status"))
        self.assertFalse(hasattr(verify, "SUCCESS_WORDS"))

    def test_other_checks_still_work(self):
        self.assertTrue(verify.run_checks(shot_text(), "412345678901", 4000))       # wrong amount
        self.assertTrue(verify.run_checks(shot_text(), "999999999999", 5000))       # id not on screenshot
        self.assertTrue(verify.run_checks(shot_text(), "12345", 5000))              # too short

    def test_receiver_optional(self):
        self.assertEqual(verify.run_checks(shot_text(), "412345678901", 5000, ""), [])
        self.assertEqual(verify.run_checks(shot_text(), "412345678901", 5000, "shop@okbank"), [])
        self.assertTrue(verify.run_checks(shot_text(), "412345678901", 5000, "other@ybl"))


class StoreCase(unittest.TestCase):
    """Base: self.xl is the 'target' handed to store functions (an Excel path or a backend)."""
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.make_target()

    def make_target(self):
        self.xl = Path(self.dir) / "payments.xlsx"

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def online(self, txn="412345678901", h="hash1", amount=5000, name="Hardik", plot="87",
               link="https://drive.google.com/file/d/abc/view"):
        return store.add_online_payment(name, plot, amount, txn, h, receipt_link=link, file_path=self.xl)


class StoreTests:
    def test_transaction_id_is_receipt_no(self):
        rec = self.online()
        self.assertEqual(rec["Receipt No"], "412345678901")
        df = store.load_payments(self.xl)
        self.assertEqual(df.iloc[0]["Receipt No"], "412345678901")
        self.assertEqual(df.iloc[0]["Source"], "Online")
        self.assertEqual(df.iloc[0]["Bank Status"], "Unconfirmed")
        self.assertEqual(df.iloc[0]["Payment Mode"], "Online")
        self.assertEqual(df.iloc[0]["Receipt Link"], "https://drive.google.com/file/d/abc/view")
        self.assertNotIn("Screenshot File", df.columns)

    def test_long_txn_id_kept_as_text(self):
        self.online(txn="T2610311042123456789012345")
        self.assertEqual(store.load_payments(self.xl).iloc[0]["Receipt No"], "T2610311042123456789012345")

    def test_duplicate_txn_and_lookalike(self):
        self.online()
        with self.assertRaises(store.DuplicateError):
            self.online(h="other")
        with self.assertRaises(store.DuplicateError):       # same id, spaces/dashes differ
            self.online(txn="4123 4567-8901", h="other2")
        self.assertEqual(len(store.load_payments(self.xl)), 1)

    def test_duplicate_screenshot(self):
        self.online()
        with self.assertRaises(store.DuplicateError):
            self.online(txn="999999999999", h="hash1")

    def test_txn_cannot_reuse_cash_receipt_no(self):
        store.add_admin_payment({"Receipt No": "RCP-0001", "Date": "x", "Name": "a", "Plot No": "1",
                                 "Amount": "10.00", "Payment Mode": "Cash"}, self.xl)
        with self.assertRaises(store.DuplicateError):
            self.online(txn="RCP-0001")

    def test_receipt_sequence_ignores_transaction_ids(self):
        self.assertEqual(store.next_receipt_no(self.xl), "RCP-0001")
        store.add_admin_payment({"Receipt No": "RCP-0001", "Name": "a"}, self.xl)
        self.online(txn="412345678901")                     # huge number must not leak into RCP-
        self.assertEqual(store.next_receipt_no(self.xl), "RCP-0002")

    def test_rejected_attempt_log(self):
        self.assertTrue(store.load_attempts(self.xl).empty)
        store.log_attempt("n", "1", 5, "UTR1", "Amount mismatch", self.xl)
        att = store.load_attempts(self.xl)
        self.assertEqual(att.iloc[0]["Reason"], "Amount mismatch")
        self.assertEqual(len(store.load_payments(self.xl)), 0)   # not mixed with payments

    def test_concurrent_submissions_none_lost(self):
        store._ensure_workbook(self.xl)
        errs = []

        def work(i):
            try:
                self.online(txn=f"UTR{i:010d}", h=f"h{i}")
            except Exception as e:  # pragma: no cover
                errs.append(e)

        ts = [threading.Thread(target=work, args=(i,)) for i in range(12)]
        [t.start() for t in ts]
        [t.join() for t in ts]
        self.assertEqual(errs, [])
        self.assertEqual(len(store.load_payments(self.xl)), 12)



class ExcelTarget:
    def make_target(self):
        self.xl = Path(self.dir) / "payments.xlsx"


class SheetsTarget:
    def make_target(self):
        self.fake = FakeSheetsService()
        drive_uploader.get_sheets_service = lambda: self.fake
        self.xl = store.SheetsBackend(spreadsheet_id="fake-id")


class BankCheckTests:
    def test_match_and_confirm(self):
        self.online(txn="412345678901", h="a", amount=5000)
        self.online(txn="509876543210", h="b", amount=1500)
        self.online(txn="777777777777", h="c", amount=900)
        store.add_admin_payment({"Receipt No": "RCP-0001", "Amount": "5000.00", "Name": "cash"}, self.xl)

        stmt = pd.DataFrame({
            "Date": ["01-11-2026", "01-11-2026", "02-11-2026"],
            "Narration": ["UPI/412345678901/hardik", "UPI/509876543210/ramesh", "UPI/777777777777/x"],
            "Credit": ["5,000.00", "1,499.00", "900.00"],        # middle amount differs
        })
        online = store.load_payments(self.xl)
        online = online[online["Source"] == "Online"].to_dict("records")
        found = bank_check.match_statement(stmt, online)
        self.assertEqual(sorted(found), ["412345678901", "777777777777"])

        self.assertEqual(store.mark_bank_confirmed(found, self.xl), 2)
        df = store.load_payments(self.xl).set_index("Receipt No")
        self.assertEqual(df.loc["412345678901", "Bank Status"], "Confirmed")
        self.assertEqual(df.loc["509876543210", "Bank Status"], "Unconfirmed")
        self.assertEqual(df.loc["RCP-0001", "Bank Status"], "")  # cash rows untouched
        self.assertEqual(store.mark_bank_confirmed(found, self.xl), 0)   # idempotent


class MoreStoreTests:
    def test_update_receipt_link_and_text_preserved(self):
        store.add_admin_payment({"Receipt No": "RCP-0001", "Name": "a", "Amount": "5000.00",
                                 "Phone": "0098765"}, self.xl)
        store.update_receipt_link("RCP-0001", "https://drive/pdf", self.xl)
        row = store.load_payments(self.xl).iloc[0]
        self.assertEqual((row["Receipt Link"], row["Phone"], row["Amount"], row["Source"]),
                         ("https://drive/pdf", "0098765", "5000.00", "Admin"))   # leading zeros kept

    def test_online_without_drive_link_is_still_saved(self):
        rec = store.add_online_payment("N", "1", 10, "UTR0000000001", "hh", remarks="Screenshot could not be saved to Drive", file_path=self.xl)
        row = store.load_payments(self.xl).iloc[0]
        self.assertEqual((row["Receipt Link"], row["Remarks"]), ("", "Screenshot could not be saved to Drive"))


class TestStoreExcel(ExcelTarget, StoreTests, MoreStoreTests, BankCheckTests, StoreCase):
    pass


class TestStoreSheets(SheetsTarget, StoreTests, MoreStoreTests, BankCheckTests, StoreCase):
    def test_everything_is_written_as_text(self):
        self.online()
        # (the fake asserts RAW + str on every append); first tab has the headers
        self.assertEqual(self.fake.tabs["Payments"][0], store.HEADERS)
        self.assertEqual(self.fake.tabs["Payments"][1][0], "412345678901")

    def test_old_sheet_with_10_columns_is_migrated(self):
        old = ["Receipt No", "Date", "Name", "Plot No", "Amount", "Payment Mode",
               "Email", "Phone", "Remarks", "Receipt Link"]
        self.fake.tabs["Payments"] = [old,
            ["RCP-0001", "d", "hardik", "87", "5000.00", "Cash", "a@b.c", "98", "", ""],
            ["RCP-0002", "d", "ramesh", "68", "1500.00", "Cash", "a@b.c", "72", "", "https://drive/x"]]
        df = store.load_payments(self.xl)
        self.assertEqual(list(df.columns), store.HEADERS)
        self.assertEqual(df["Source"].tolist(), ["Admin", "Admin"])
        self.assertEqual(store.next_receipt_no(self.xl), "RCP-0003")

    def test_new_empty_spreadsheet_gets_headers(self):
        self.assertEqual(store.next_receipt_no(self.xl), "RCP-0001")
        self.assertEqual(self.fake.tabs["Payments"][0], store.HEADERS)

    def test_duplicate_check_ignores_read_cache(self):
        self.online()
        store.load_payments(self.xl)                  # warms the 4-second read cache
        with self.assertRaises(store.DuplicateError):
            self.online(h="other")                    # must still see the row just written

    def test_admin_reads_are_cached_to_save_quota(self):
        store.load_payments(self.xl)
        before = self.fake.calls
        for _ in range(5):
            store.load_payments(self.xl)
        self.assertEqual(self.fake.calls, before)


class TestExcelOnly(ExcelTarget, StoreCase):
    def test_column_order_independent(self):
        store._ensure_workbook(self.xl)
        wb = load_workbook(self.xl)
        ws = wb["Payments"]
        ws.cell(1, 1).value, ws.cell(1, 3).value = "Name", "Receipt No"      # swap two headers
        wb.save(self.xl)
        store.save_payment({"Receipt No": "R1", "Name": "N"}, self.xl)
        df = store.load_payments(self.xl)
        self.assertEqual((df.iloc[0]["Receipt No"], df.iloc[0]["Name"]), ("R1", "N"))




class TestMigration(ExcelTarget, StoreCase):
    def test_old_admin_file_is_migrated_and_preserved(self):
        old = ["Receipt No", "Date", "Name", "Plot No", "Amount", "Payment Mode",
               "Email", "Phone", "Remarks", "Receipt Link"]
        wb = Workbook(); ws = wb.active; ws.title = "Payments"
        ws.append(old)
        ws.append(["RCP-0001", "31-10-2026 15:30", "hardik", "87", "5000.00", "Cash", "a@b.c", "9893045455", None, None])
        ws.append(["RCP-0002", "31-10-2026 17:55", "ramesh", "68", "1500.00", "Cash", "a@b.c", "7224010897", None, "https://drive/x"])
        wb.save(self.xl)

        store.migrate_excel_if_needed(self.xl)
        df = store.load_payments(self.xl)
        self.assertEqual(list(df.columns), store.HEADERS)
        self.assertEqual(df["Receipt No"].tolist(), ["RCP-0001", "RCP-0002"])
        self.assertEqual(df["Source"].tolist(), ["Admin", "Admin"])
        self.assertEqual(df.iloc[1]["Receipt Link"], "https://drive/x")
        self.assertEqual(df.iloc[0]["Amount"], "5000.00")
        self.assertEqual(store.next_receipt_no(self.xl), "RCP-0003")
        self.online()                                           # still writable afterwards
        self.assertEqual(len(store.load_payments(self.xl)), 3)

    def test_migration_idempotent(self):
        store._ensure_workbook(self.xl)
        before = self.xl.read_bytes()
        store.migrate_excel_if_needed(self.xl)
        self.assertEqual(before, self.xl.read_bytes())          # untouched when nothing to add


class TestExportAndDrive(ExcelTarget, StoreCase):
    def test_export_has_both_tabs(self):
        self.online()
        store.log_attempt("n", "1", 5, "U", "why", self.xl)
        wb = load_workbook(io.BytesIO(store.export_xlsx_bytes(self.xl)))
        self.assertEqual(wb.sheetnames, ["Payments", "Rejected Attempts"])

    def _fake_drive(self):
        calls = {"perm": 0, "created": []}

        class Q:
            def __init__(s, v): s.v = v
            def execute(s): return s.v

        class Files:
            def list(s, q, **k):
                if "mimeType = 'application/vnd.google-apps.folder'" in q:
                    return Q({"files": [{"id": "F1"}]})
                return Q({"files": []})
            def create(s, body, media_body=None, fields=None):
                calls["created"].append(body["name"]); return Q({"id": "ID1", "webViewLink": "https://drive/ID1"})
            def get(s, **k): return Q({"webViewLink": "https://drive/ID1"})

        class Perms:
            def create(s, **k): calls["perm"] += 1; return Q({})

        class Svc:
            def files(s): return Files()
            def permissions(s): return Perms()

        drive_uploader.get_drive_service = lambda: Svc()
        drive_uploader._FOLDER_CACHE.clear()
        drive_uploader.MediaIoBaseUpload = lambda *a, **k: object()
        return calls

    def test_screenshot_upload_private_by_default(self):
        calls = self._fake_drive()
        link = drive_uploader.upload_screenshot(b"png-bytes", "UTR1_abcd1234.png", "image/png")
        self.assertEqual(link, "https://drive/ID1")
        self.assertEqual(calls["created"], ["UTR1_abcd1234.png"])
        self.assertEqual(calls["perm"], 0)                       # not shared with the world

    def test_screenshot_upload_public_when_enabled(self):
        calls = self._fake_drive()
        import os
        os.environ["SCREENSHOT_LINK_PUBLIC"] = "true"
        try:
            drive_uploader.upload_screenshot(b"x", "a.png", "image/png")
        finally:
            del os.environ["SCREENSHOT_LINK_PUBLIC"]
        self.assertEqual(calls["perm"], 1)


class TestReceipts(unittest.TestCase):
    def test_pdf_for_transaction_id_receipt_no(self):
        rec = {"Receipt No": "UTR/412 345:6789", "Date": "31-10-2026 10:00", "Name": "A", "Plot No": "1",
               "Amount": "5000.00", "Payment Mode": "Online", "Remarks": ""}
        p = receipts.generate_receipt_pdf(rec)
        try:
            self.assertTrue(p.exists() and p.name == "UTR_412_345_6789.pdf")
        finally:
            p.unlink(missing_ok=True)

    def test_amount_words(self):
        self.assertEqual(receipts.amount_in_words("125000.50"),
                         "Rupees One Lakh Twenty Five Thousand and Fifty Paise Only")


if __name__ == "__main__":
    unittest.main(verbosity=2)
